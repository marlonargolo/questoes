"""Importa bases já estruturadas (CSV, XLSX, JSON, JSONL, Parquet) com um mapa de campos.

Serve para datasets abertos (Hugging Face, Kaggle), bases compradas/licenciadas
ou exportações de outros sistemas. O mapa diz de qual campo da origem vem cada coluna.

Exemplo de mapa (YAML):
    fixos: {Banca: FGV, Orgao: OAB}
    campos:
      Ano: exam_year
      Enunciado: question
      Gabarito: answerKey
    alternativas: choices.text        # lista; ou Alternativa_A..E em "campos"
    certo_errado: false
"""
import csv
import json
from pathlib import Path

import yaml

from ..schema import COLUNAS, Questao


def _get(registro, caminho):
    """Lê 'a.b.c' em dicts aninhados; aceita índice numérico em listas."""
    atual = registro
    for parte in str(caminho).split("."):
        if isinstance(atual, dict):
            atual = atual.get(parte)
        elif isinstance(atual, list) and parte.isdigit():
            atual = atual[int(parte)] if int(parte) < len(atual) else None
        else:
            return None
        if atual is None:
            return None
    return atual


def ler_registros(caminho):
    caminho = Path(caminho)
    ext = caminho.suffix.lower()
    if ext == ".jsonl":
        with open(caminho, encoding="utf-8") as f:
            for linha in f:
                if linha.strip():
                    yield json.loads(linha)
    elif ext == ".json":
        with open(caminho, encoding="utf-8") as f:
            dados = json.load(f)
        yield from (dados if isinstance(dados, list) else dados.get("data", []))
    elif ext in (".csv", ".tsv"):
        with open(caminho, encoding="utf-8-sig", newline="") as f:
            amostra = f.read(4096)
            f.seek(0)
            dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
            yield from csv.DictReader(f, dialect=dialeto)
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(caminho, read_only=True).active
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c) if c is not None else "" for c in next(linhas)]
        for l in linhas:
            yield {k: ("" if v is None else v) for k, v in zip(cab, l)}
    elif ext == ".parquet":
        import pyarrow.parquet as pq

        yield from pq.read_table(caminho).to_pylist()
    else:
        raise ValueError(f"extensão não suportada: {ext}")


def carregar_mapa(caminho_ou_dict):
    if isinstance(caminho_ou_dict, dict):
        return caminho_ou_dict
    with open(caminho_ou_dict, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def importar(caminho, mapa=None):
    """Gera Questao a partir de um arquivo. Sem mapa, assume as 13 colunas com os nomes exatos."""
    mapa = carregar_mapa(mapa) if mapa else {"campos": {c: c for c in COLUNAS}}
    fixos = mapa.get("fixos", {})
    campos = mapa.get("campos", {})
    alt_campo = mapa.get("alternativas")
    ce = mapa.get("certo_errado", False)
    filtro = mapa.get("ignorar_se_verdadeiro")  # ex.: "nullified" (questões anuladas)

    for reg in ler_registros(caminho):
        if filtro and _get(reg, filtro) in (True, "true", "True", 1, "1"):
            continue
        d = {c: "" for c in COLUNAS}
        d.update({k: str(v) for k, v in fixos.items()})
        for col, origem in campos.items():
            v = _get(reg, origem)
            if v is not None:
                d[col] = str(v)
        if alt_campo:
            alts = _get(reg, alt_campo) or []
            if isinstance(alts, str):
                alts = json.loads(alts)
            for letra, texto in zip("ABCDE", alts):
                d[f"Alternativa_{letra}"] = str(texto)
        if ce:
            gab = d.pop("Gabarito")
            for l in "ABCDE":
                d.pop(f"Alternativa_{l}")
            yield Questao.certo_errado(gab, **d)
        else:
            yield Questao(**d)
