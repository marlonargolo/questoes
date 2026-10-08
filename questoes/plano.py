"""Plano de coleta: planilha de provas por banca/ano -> resumo e manifesto YAML.

Fluxo: preencha Arquivo_Prova, Arquivo_Gabarito e Faixas_Disciplinas de cada
linha, mude Status para "pronto" e gere o manifesto com `python -m questoes plano`.
Faixas_Disciplinas: "1-13:Língua Portuguesa | 14-25:Direito Administrativo".
"""
import csv
from collections import defaultdict
from pathlib import Path

import yaml

COLUNAS_PLANO = ["ID", "Banca", "Orgao", "Ano", "Cargo_Prova", "Tipo", "Alternativas", "Questoes_Estimadas",
                 "Confianca", "Status", "Arquivo_Prova", "Arquivo_Gabarito", "Faixas_Disciplinas", "Observacoes"]


def ler_plano(caminho):
    caminho = Path(caminho)
    if caminho.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(caminho, read_only=True)["Plano"]
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c) for c in next(linhas)]
        return [{k: "" if v is None else str(v) for k, v in zip(cab, l)} for l in linhas if any(l)]
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def parse_faixas(texto):
    faixas = []
    for parte in filter(None, (p.strip() for p in (texto or "").split("|"))):
        intervalo, disciplina = parte.split(":", 1)
        de, _, ate = intervalo.partition("-")
        faixas.append({"de": int(de), "ate": int(ate or de), "disciplina": disciplina.strip()})
    return faixas


def resumo(linhas):
    por_banca = defaultdict(lambda: [0, 0])
    por_banca_ano = defaultdict(lambda: [0, 0])
    for l in linhas:
        q = int(l["Questoes_Estimadas"] or 0)
        for chave, d in (((l["Banca"],), por_banca), ((l["Banca"], l["Ano"]), por_banca_ano)):
            d[chave][0] += 1
            d[chave][1] += q
    return por_banca, por_banca_ano


def gerar_manifesto(linhas, saida):
    """Gera manifesto só das linhas com Status=pronto. Retorna (gerados, pendencias)."""
    provas, pend = [], []
    for l in linhas:
        if l["Status"].strip().lower() != "pronto":
            continue
        falta = [c for c in ("Arquivo_Prova", "Arquivo_Gabarito") if not l[c].strip()]
        if l["Tipo"] not in ("certo_errado", "multipla_escolha"):
            falta.append("Tipo (defina certo_errado ou multipla_escolha)")
        if falta:
            pend.append(f"ID {l['ID']} ({l['Orgao']} {l['Ano']} {l['Cargo_Prova']}): falta {', '.join(falta)}")
            continue
        item = {"prova": l["Arquivo_Prova"], "gabarito": l["Arquivo_Gabarito"], "banca": l["Banca"],
                "orgao": l["Orgao"], "ano": int(l["Ano"]), "tipo": l["Tipo"]}
        if l["Tipo"] == "multipla_escolha":
            item["numero_alternativas"] = int(l["Alternativas"] or 5)
        faixas = parse_faixas(l["Faixas_Disciplinas"])
        if faixas:
            item["disciplinas"] = faixas
        provas.append(item)
    with open(saida, "w", encoding="utf-8") as f:
        yaml.safe_dump({"padrao": {"colunas": 2, "paginas_ignorar": [1]}, "provas": provas}, f,
                       allow_unicode=True, sort_keys=False)
    return len(provas), pend


def exportar_xlsx(linhas, saida):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Plano"
    ws.append(COLUNAS_PLANO)
    for l in linhas:
        ws.append([int(l[c]) if c in ("ID", "Ano", "Alternativas", "Questoes_Estimadas") and str(l[c]).isdigit()
                   else l[c] for c in COLUNAS_PLANO])
    for c in ws[1]:
        c.font = Font(bold=True)
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        if row[8].value == "verificar":
            row[8].fill = amarelo
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, larg in zip("ABCDEFGHIJKLMN", [5, 11, 34, 6, 40, 16, 6, 10, 10, 10, 30, 30, 40, 45]):
        ws.column_dimensions[col].width = larg
    dv = DataValidation(type="list", formula1='"a_coletar,baixado,pronto,importado,descartado"')
    ws.add_data_validation(dv)
    dv.add(f"J2:J{ws.max_row}")

    pb, pba = resumo(linhas)
    r = wb.create_sheet("Resumo")
    r.append(["Banca", "Provas", "Questões estimadas"])
    for (b,), (n, q) in sorted(pb.items()):
        r.append([b, n, q])
    r.append(["TOTAL", sum(v[0] for v in pb.values()), sum(v[1] for v in pb.values())])
    r.append([])
    r.append(["Banca", "Ano", "Provas", "Questões estimadas"])
    for (b, a), (n, q) in sorted(pba.items()):
        r.append([b, int(a), n, q])
    for row in r.iter_rows():
        if row[0].value in ("Banca", "TOTAL"):
            for c in row:
                c.font = Font(bold=True)
    for col, larg in zip("ABCD", [14, 10, 10, 18]):
        r.column_dimensions[col].width = larg
    wb.save(saida)
