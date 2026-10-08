import csv
import json

from openpyxl import load_workbook

from questoes.pipeline import consolidar
from questoes.schema import COLUNAS


def test_consolidar_exporta_13_colunas(tmp_path):
    entrada = tmp_path / "e.jsonl"
    linhas = [
        dict(Banca="FCC", Orgao="TRT", Ano="2023", Disciplina="Língua Portuguesa", Enunciado="<p>Há ocorrência de crase em:</p>",
             Alternativa_A="(A) a", Alternativa_B="(B) à noite", Alternativa_C="(C) c", Alternativa_D="(D) d",
             Alternativa_E="(E) e", Gabarito=" b"),
        # duplicada com HTML diferente
        dict(Banca="FCC", Orgao="TRT", Ano="2023", Disciplina="Língua Portuguesa", Enunciado="Há ocorrência de crase em:",
             Alternativa_A="a", Alternativa_B="à noite", Alternativa_C="c", Alternativa_D="d", Alternativa_E="e",
             Gabarito="B", Comentario_Professor="Crase antes de locução adverbial feminina."),
        dict(Banca="", Orgao="X", Ano="2020", Disciplina="D", Enunciado="Sem banca, deve ser rejeitada.",
             Alternativa_A="Certo", Alternativa_B="Errado", Gabarito="A"),
    ]
    entrada.write_text("\n".join(json.dumps(l, ensure_ascii=False) for l in linhas), encoding="utf-8")
    csv_out, xlsx_out = tmp_path / "b.csv", tmp_path / "b.xlsx"
    rel = consolidar([entrada], [csv_out, xlsx_out], rejeitadas=tmp_path / "r.csv")
    assert (rel["validas"], rel["rejeitadas"], rel["duplicadas_removidas"]) == (1, 1, 1)

    with open(csv_out, encoding="utf-8", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    assert linhas[0] == COLUNAS
    linha = dict(zip(COLUNAS, linhas[1]))
    assert linha["Enunciado"] == "Há ocorrência de crase em:"
    assert linha["Alternativa_B"] == "à noite"
    assert linha["Gabarito"] == "B"
    assert linha["Assunto"] == "Crase"
    assert linha["Comentario_Professor"].startswith("Crase antes")

    ws = load_workbook(xlsx_out).active
    assert [c.value for c in ws[1]] == COLUNAS
    assert ws.cell(2, 3).value == 2023
